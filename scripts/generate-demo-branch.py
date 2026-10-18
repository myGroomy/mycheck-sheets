#!/usr/bin/env python3
"""Generate a populated flat-mode branch workbook for local testing."""

from __future__ import annotations

import argparse
import hashlib
import json
from datetime import date, datetime, time, timedelta, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


STATIC_SHEETS = {
    "ShiftDefinitions": ["id", "name", "start_time", "end_time", "crosses_midnight", "sort_order", "is_active", "created_at", "updated_at", "version"],
    "SopCategories": ["id", "shift_definition_id", "name", "sort_order", "is_active", "created_at", "updated_at", "version"],
    "ChecklistPoints": ["id", "sop_category_id", "title", "instruction", "input_type", "input_types", "options", "is_required", "target_time", "tolerance_minutes", "active_days", "number_min", "number_max", "sort_order", "is_active", "created_at", "updated_at", "version"],
    "HandoverFields": ["id", "shift_definition_id", "label", "field_type", "input_types", "options", "is_required", "sort_order", "created_at", "updated_at", "version"],
    "ShiftInstances": ["id", "shift_definition_id", "shift_date", "tab_month", "status", "pj_user_id", "opened_by", "opened_at", "opened_outside_hours", "closed_at", "closed_by", "close_type", "force_close_reason", "is_incomplete", "no_incident_confirmed", "void_reason", "void_by", "void_at", "is_test", "snapshot_encoding", "template_snapshot", "snapshot_hash", "created_at", "updated_at", "version", "tab_year"],
    "Participants": ["id", "shift_instance_id", "user_id", "first_action_at", "first_action_type", "created_at"],
    "Reports": ["id", "shift_instance_id", "report_number", "generated_by", "generated_at", "is_locked", "summary_stats", "content_hash", "unlock_count", "last_unlocked_at", "last_unlocked_by", "created_at", "updated_at", "version"],
    "Addenda": ["id", "report_id", "author_id", "note", "created_at"],
    "Summary": ["id", "summary_date", "shift_definition_id", "shifts_total", "shifts_closed_normal", "shifts_closed_forced", "shifts_void", "required_total", "required_done", "required_skipped", "timed_on_time", "timed_early", "timed_late", "incidents_total", "incidents_open", "incidents_by_category", "handovers_read", "participants_count", "computed_at"],
    "Snapshots": ["id", "shift_instance_id", "part_no", "chunk", "created_at"],
    "IncidentIndex": ["incident_id", "tab_month", "status", "category_id", "shift_instance_id", "outside_shift", "reported_at", "is_test", "updated_at"],
    "IncidentCategories": ["id", "name", "sort_order", "is_active", "created_at", "updated_at", "points"],
    "Notifications": ["id", "user_id", "type", "title", "body", "link", "is_read", "created_at"],
    "Drafts": ["id", "shift_instance_id", "user_id", "shift_date", "payload", "client_updated_at", "created_at", "updated_at", "version"],
}

PERIOD_SHEETS = {
    "Entries": ["id", "shift_instance_id", "point_ref", "state", "value", "out_of_range", "photo_ids", "completed_by", "completed_at", "timing_label", "timing_delta_minutes", "skip_reason", "created_at", "updated_at", "version"],
    "EntryLogs": ["id", "shift_instance_id", "entry_id", "point_ref", "action", "outcome", "user_id", "winner_user_id", "prev_state", "new_state", "value", "note", "client_action_id", "client_at", "at", "created_at"],
    "Handovers": ["id", "shift_instance_id", "values", "free_text", "photo_ids", "submitted_by", "submitted_at", "created_at", "updated_at"],
    "HandoverAcks": ["id", "handover_id", "reading_shift_instance_id", "user_id", "read_at", "created_at"],
    "Incidents": ["id", "shift_instance_id", "tab_month", "category_id", "description", "occurred_at", "reported_by", "reported_at", "status", "outside_shift", "link_source", "linked_by", "linked_at", "source_entry_id", "severity", "status_changed_by", "status_changed_at", "is_test", "created_at", "updated_at", "version"],
    "IncidentNotes": ["id", "incident_id", "author_id", "author_role", "note", "created_at"],
    "Photos": ["id", "shift_instance_id", "owner_type", "owner_id", "storage", "file_ref", "mime", "size_bytes", "width", "height", "sort_order", "status", "uploaded_by", "uploaded_at", "purged_at", "created_at", "sha256"],
    "AuditLog": ["id", "seq", "at", "actor_id", "action", "object_type", "object_id", "branch_id", "shift_instance_id", "before", "after", "reason", "prev_hash", "hash"],
}

USERS = [
    {"id": "USRDEMO01", "name": "Alya Pratama", "role": "petugas"},
    {"id": "USRDEMO02", "name": "Fajar Nugraha", "role": "petugas"},
    {"id": "USRDEMO03", "name": "Sinta Maharani", "role": "petugas"},
]


def iso(value: datetime) -> str:
    return value.astimezone(timezone.utc).isoformat().replace("+00:00", "Z")


def stable_id(prefix: str, value: str) -> str:
    return f"{prefix}{hashlib.sha1(value.encode()).hexdigest()[:12].upper()}"


def add_sheet(workbook: Workbook, title: str, headers: list[str]) -> None:
    sheet = workbook.create_sheet(title)
    sheet.append(headers)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
    sheet.row_dimensions[1].height = 30
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for index, header in enumerate(headers, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = min(max(len(header) + 2, 14), 30)


def append_row(workbook: Workbook, sheet_name: str, headers: list[str], row: dict[str, object]) -> None:
    workbook[sheet_name].append([row.get(header, "") for header in headers])


def snapshot_for(shift_id: str, shift_name: str, categories: list[tuple[str, str, list[str]]], handover_fields: list[tuple[str, str]], start_time: str, end_time: str) -> dict[str, object]:
    return {
        "v": 1,
        "shift": {"id": shift_id, "name": shift_name, "start_time": start_time, "end_time": end_time, "crosses_midnight": False},
        "settings": {"tolerance_default_minutes": 15, "timezone": "Asia/Jakarta"},
        "categories": [{
            "id": category_id,
            "name": category_name,
            "sort_order": category_index,
            "points": [{
                "point_ref": point_id,
                "title": f"{category_name} - item {point_index}",
                "instruction": "Ikuti SOP dan simpan bukti bila diperlukan.",
                "input_type": "ok_tidak_ok",
                "input_types": ["ok_tidak_ok"],
                "options": None,
                "is_required": True,
                "target_time": None,
                "tolerance_minutes": 15,
                "active_days": None,
                "number_min": None,
                "number_max": None,
                "sort_order": point_index,
            } for point_index, point_id in enumerate(point_ids, start=1)],
        } for category_index, (category_id, category_name, point_ids) in enumerate(categories, start=1)],
        "handover_fields": [{"id": field_id, "label": field_label, "field_type": "teks", "input_types": ["teks"], "options": None, "is_required": True, "sort_order": index} for index, (field_id, field_label) in enumerate(handover_fields, start=1)],
    }


def generate_demo(output: Path, start_month: date) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    add_sheet(workbook, "_meta", ["key", "value"])
    meta = workbook["_meta"]
    created_at = iso(datetime.now(timezone.utc))
    for key, value in [("schema_version", "2"), ("period_mode", "flat"), ("branch_id", "CBDG-DEMO"), ("created_at", created_at), ("tolerance_default_minutes", "15"), ("version", "2")]:
        meta.append([key, value])
    meta.column_dimensions["A"].width = 32
    meta.column_dimensions["B"].width = 42

    for title, headers in STATIC_SHEETS.items():
        add_sheet(workbook, title, headers)
    for title, headers in PERIOD_SHEETS.items():
        add_sheet(workbook, title, headers)

    now = datetime.now(timezone.utc)
    shift_defs = [
        ("SHIFT-PAGI", "Shift Pagi", "07:00", "15:00"),
        ("SHIFT-MALAM", "Shift Malam", "15:00", "23:00"),
    ]
    category_by_shift: dict[str, list[tuple[str, str, list[str]]]] = {}
    sop_names = ["Kebersihan", "Persiapan Operasional", "Peralatan", "Stok dan Display", "Pelayanan"]
    for shift_id, shift_name, start_time, end_time in shift_defs:
        categories: list[tuple[str, str, list[str]]] = []
        for category_index, category_name in enumerate(sop_names, start=1):
            category_id = f"SOP-{shift_id}-{category_index}"
            point_ids = [f"{category_id}-P{point_index}" for point_index in range(1, 6)]
            categories.append((category_id, category_name, point_ids))
        category_by_shift[shift_id] = categories
        append_row(workbook, "ShiftDefinitions", STATIC_SHEETS["ShiftDefinitions"], {
            "id": shift_id, "name": shift_name, "start_time": start_time, "end_time": end_time,
            "crosses_midnight": False, "sort_order": 1 if shift_id.endswith("PAGI") else 2,
            "is_active": True, "created_at": created_at, "updated_at": created_at, "version": 1,
        })
        for category_index, (category_id, category_name, point_ids) in enumerate(categories, start=1):
            append_row(workbook, "SopCategories", STATIC_SHEETS["SopCategories"], {
                "id": category_id, "shift_definition_id": shift_id, "name": category_name,
                "sort_order": category_index, "is_active": True, "created_at": created_at, "updated_at": created_at, "version": 1,
            })
            for point_index, point_id in enumerate(point_ids, start=1):
                append_row(workbook, "ChecklistPoints", STATIC_SHEETS["ChecklistPoints"], {
                    "id": point_id, "sop_category_id": category_id, "title": f"{category_name} - item {point_index}",
                    "instruction": "Ikuti SOP dan pastikan kondisi sesuai standar.", "input_type": "ok_tidak_ok",
                    "input_types": json.dumps(["ok_tidak_ok"]), "options": "", "is_required": True,
                    "target_time": "08:00" if shift_id.endswith("PAGI") else "16:00", "tolerance_minutes": 15,
                    "active_days": "", "number_min": "", "number_max": "", "sort_order": point_index,
                    "is_active": True, "created_at": created_at, "updated_at": created_at, "version": 1,
                })
        handover_fields = [(f"HVR-{shift_id}-{index}", label) for index, label in enumerate(["Kondisi outlet", "Catatan operasional", "Stok kritis", "Kendala peralatan", "Arahan shift berikutnya"], start=1)]
        for field_index, (field_id, field_label) in enumerate(handover_fields, start=1):
            append_row(workbook, "HandoverFields", STATIC_SHEETS["HandoverFields"], {
                "id": field_id, "shift_definition_id": shift_id, "label": field_label,
                "field_type": "teks", "input_types": json.dumps(["teks"]), "options": "", "is_required": True,
                "sort_order": field_index, "created_at": created_at, "updated_at": created_at, "version": 1,
            })

    incident_categories = [
        ("INC-CLEAN", "Kebersihan belum sesuai", 5),
        ("INC-STOCK", "Stok atau peralatan bermasalah", 8),
        ("INC-SERVICE", "Kendala pelayanan", 10),
    ]
    for index, (category_id, name, points) in enumerate(incident_categories, start=1):
        append_row(workbook, "IncidentCategories", STATIC_SHEETS["IncidentCategories"], {
            "id": category_id, "name": name, "sort_order": index, "is_active": True,
            "created_at": created_at, "updated_at": created_at, "points": points,
        })

    sequence = 0
    summary_by_day: dict[str, dict[str, int]] = {}
    for month_offset in range(3):
        month_start = date(start_month.year + (start_month.month - 1 + month_offset) // 12, (start_month.month - 1 + month_offset) % 12 + 1, 1)
        for day_number in range(1, 8):
            shift_date = month_start + timedelta(days=day_number - 1)
            day_key = shift_date.isoformat()
            summary_by_day.setdefault(day_key, {"total": 0, "done": 0, "skipped": 0, "incidents": 0})
            for shift_index, (shift_id, shift_name, start_time, end_time) in enumerate(shift_defs):
                sequence += 1
                shift_instance_id = f"SHIFT-{shift_date.strftime('%Y%m%d')}-{shift_index + 1}"
                pj = USERS[(sequence - 1) % len(USERS)]
                participant_ids = [USERS[(sequence - 1) % 3]["id"], USERS[sequence % 3]["id"]]
                opened_local = datetime.combine(shift_date, time.fromisoformat(start_time))
                closed_local = datetime.combine(shift_date, time.fromisoformat(end_time))
                opened = opened_local.replace(tzinfo=timezone.utc) - timedelta(hours=7)
                closed = closed_local.replace(tzinfo=timezone.utc) - timedelta(hours=7)
                categories = category_by_shift[shift_id]
                point_ids = [point_id for _, _, category_points in categories for point_id in category_points]
                handover_fields = [(f"HVR-{shift_id}-{index}", label) for index, label in enumerate(["Kondisi outlet", "Catatan operasional", "Stok kritis", "Kendala peralatan", "Arahan shift berikutnya"], start=1)]
                snapshot = snapshot_for(shift_id, shift_name, categories, handover_fields, start_time, end_time)
                snapshot_json = json.dumps(snapshot, separators=(",", ":"))
                snapshot_hash = hashlib.sha256(snapshot_json.encode()).hexdigest()
                append_row(workbook, "ShiftInstances", STATIC_SHEETS["ShiftInstances"], {
                    "id": shift_instance_id, "shift_definition_id": shift_id, "shift_date": day_key,
                    "tab_month": day_key[:7], "status": "ditutup", "pj_user_id": pj["id"], "opened_by": pj["id"],
                    "opened_at": iso(opened), "opened_outside_hours": False, "closed_at": iso(closed), "closed_by": pj["id"],
                    "close_type": "normal", "force_close_reason": "", "is_incomplete": False, "no_incident_confirmed": True,
                    "void_reason": "", "void_by": "", "void_at": "", "is_test": False, "snapshot_encoding": "json",
                    "template_snapshot": snapshot_json, "snapshot_hash": snapshot_hash, "created_at": iso(opened), "updated_at": iso(closed),
                    "version": 1, "tab_year": day_key[:4],
                })
                for participant_index, user_id in enumerate(participant_ids):
                    joined_at = opened + timedelta(minutes=participant_index * 3)
                    append_row(workbook, "Participants", STATIC_SHEETS["Participants"], {
                        "id": f"PART-{sequence:03d}-{participant_index + 1}", "shift_instance_id": shift_instance_id,
                        "user_id": user_id, "first_action_at": iso(joined_at),
                        "first_action_type": "saya_bertugas" if participant_index == 0 else "buka_shift", "created_at": iso(joined_at),
                    })
                done_count = 0
                skipped_count = 0
                for point_index, point_id in enumerate(point_ids):
                    entry_id = f"ENTRY-{sequence:03d}-{point_index + 1}"
                    worker = USERS[(point_index + sequence) % len(USERS)]
                    skipped = sequence % 5 == 0 and point_index == 3
                    state = "skip" if skipped else "selesai"
                    completed_at = closed - timedelta(minutes=30 + point_index * 7)
                    if skipped:
                        skipped_count += 1
                    else:
                        done_count += 1
                    value = "" if skipped else "true"
                    append_row(workbook, "Entries", PERIOD_SHEETS["Entries"], {
                        "id": entry_id, "shift_instance_id": shift_instance_id, "point_ref": point_id, "state": state,
                        "value": value, "out_of_range": False, "photo_ids": "", "completed_by": "" if skipped else worker["id"],
                        "completed_at": "" if skipped else iso(completed_at), "timing_label": "tepat_waktu" if not skipped else "",
                        "timing_delta_minutes": 0 if not skipped else "", "skip_reason": "Peralatan sedang maintenance" if skipped else "",
                        "created_at": iso(opened), "updated_at": iso(completed_at), "version": 1,
                    })
                    append_row(workbook, "EntryLogs", PERIOD_SHEETS["EntryLogs"], {
                        "id": f"LOG-{sequence:03d}-{point_index + 1}", "shift_instance_id": shift_instance_id, "entry_id": entry_id,
                        "point_ref": point_id, "action": "skip" if skipped else "selesai", "outcome": "diterima",
                        "user_id": worker["id"], "winner_user_id": "" if skipped else worker["id"], "prev_state": "belum",
                        "new_state": state, "value": value, "note": "", "client_action_id": f"ACTION-{sequence:03d}-{point_index + 1}",
                        "client_at": iso(completed_at), "at": iso(completed_at), "created_at": iso(completed_at),
                    })
                handover_id = f"HAND-{sequence:03d}"
                append_row(workbook, "Handovers", PERIOD_SHEETS["Handovers"], {
                    "id": handover_id, "shift_instance_id": shift_instance_id,
                    "values": json.dumps({field_id: {"teks": f"{field_label}: kondisi normal pada {day_key}."} for field_id, field_label in handover_fields}),
                    "free_text": "Tidak ada catatan kritis.", "photo_ids": "", "submitted_by": pj["id"],
                    "submitted_at": iso(closed - timedelta(minutes=5)), "created_at": iso(closed - timedelta(minutes=5)), "updated_at": iso(closed),
                })
                report_id = f"REPORT-{sequence:03d}"
                summary = {"total_required": len(point_ids), "done": done_count, "skipped": skipped_count, "total_points": len(point_ids), "earned_points": done_count, "incident_penalty": 0, "score_percent": round(done_count / len(point_ids) * 100, 2)}
                content_hash = hashlib.sha256(f"{shift_instance_id}:{report_id}".encode()).hexdigest()
                append_row(workbook, "Reports", STATIC_SHEETS["Reports"], {
                    "id": report_id, "shift_instance_id": shift_instance_id, "report_number": f"R-{day_key.replace('-', '')}-{sequence:02d}",
                    "generated_by": pj["id"], "generated_at": iso(closed), "is_locked": True, "summary_stats": json.dumps(summary),
                    "content_hash": content_hash, "unlock_count": 0, "last_unlocked_at": "", "last_unlocked_by": "",
                    "created_at": iso(closed), "updated_at": iso(closed), "version": 1,
                })
                if sequence % 4 == 0:
                    incident_id = f"INC-{sequence:03d}"
                    incident_category = incident_categories[sequence % len(incident_categories)]
                    incident_time = closed - timedelta(hours=2)
                    append_row(workbook, "Incidents", PERIOD_SHEETS["Incidents"], {
                        "id": incident_id, "shift_instance_id": shift_instance_id, "tab_month": day_key[:7],
                        "category_id": incident_category[0], "description": f"Contoh incident {incident_category[1].lower()}.",
                        "occurred_at": iso(incident_time), "reported_by": participant_ids[1], "reported_at": iso(incident_time),
                        "status": "resolved", "outside_shift": False, "link_source": "otomatis", "linked_by": "", "linked_at": "",
                        "source_entry_id": "", "severity": "sedang", "status_changed_by": pj["id"], "status_changed_at": iso(closed),
                        "is_test": False, "created_at": iso(incident_time), "updated_at": iso(closed), "version": 1,
                    })
                    append_row(workbook, "IncidentIndex", STATIC_SHEETS["IncidentIndex"], {
                        "incident_id": incident_id, "tab_month": day_key[:7], "status": "resolved", "category_id": incident_category[0],
                        "shift_instance_id": shift_instance_id, "outside_shift": False, "reported_at": iso(incident_time), "is_test": False, "updated_at": iso(closed),
                    })
                    append_row(workbook, "IncidentNotes", PERIOD_SHEETS["IncidentNotes"], {
                        "id": f"NOTE-{sequence:03d}", "incident_id": incident_id, "author_id": pj["id"], "author_role": "petugas",
                        "note": "Tindakan korektif sudah dilakukan.", "created_at": iso(closed),
                    })
                    summary_by_day[day_key]["incidents"] += 1
                summary_by_day[day_key]["total"] += 1
                summary_by_day[day_key]["done"] += done_count
                summary_by_day[day_key]["skipped"] += skipped_count
                append_row(workbook, "AuditLog", PERIOD_SHEETS["AuditLog"], {
                    "id": f"AUDIT-{sequence:03d}", "seq": sequence, "at": iso(closed), "actor_id": pj["id"],
                    "action": "close_shift", "object_type": "shift_instance", "object_id": shift_instance_id,
                    "branch_id": "CBDG-DEMO", "shift_instance_id": shift_instance_id, "before": "{\"status\":\"berjalan\"}",
                    "after": "{\"status\":\"ditutup\"}", "reason": "Demo data", "prev_hash": "", "hash": hashlib.sha256(f"AUDIT-{sequence:03d}".encode()).hexdigest(),
                })

    for day_key, values in summary_by_day.items():
        for shift_id, _, _, _ in shift_defs:
            shift_count = values["total"] // 2
            append_row(workbook, "Summary", STATIC_SHEETS["Summary"], {
                "id": f"SUMMARY-{day_key}-{shift_id}", "summary_date": day_key, "shift_definition_id": shift_id,
                "shifts_total": 1, "shifts_closed_normal": 1, "shifts_closed_forced": 0, "shifts_void": 0,
                "required_total": 25, "required_done": 25 if values["done"] >= 25 else values["done"], "required_skipped": 0,
                "timed_on_time": 25, "timed_early": 0, "timed_late": 0, "incidents_total": values["incidents"] // 2,
                "incidents_open": 0, "incidents_by_category": "{}", "handovers_read": 0, "participants_count": 2, "computed_at": created_at,
            })

    for sheet in workbook.worksheets:
        sheet.sheet_view.showGridLines = False
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    print(f"created {output} ({len(workbook.sheetnames)} sheets, 42 closed shifts, flat mode)")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, default=Path("templates/mycheck-demo-branch-3months.xlsx"))
    parser.add_argument("--start-month", default="2026-07", help="Bulan awal YYYY-MM; tujuh hari pertama tiap bulan dipakai")
    args = parser.parse_args()
    year, month = (int(value) for value in args.start_month.split("-", 1))
    generate_demo(args.output, date(year, month, 1))


if __name__ == "__main__":
    main()
