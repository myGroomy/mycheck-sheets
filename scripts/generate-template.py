#!/usr/bin/env python3
"""Generate an empty MyCheck branch template for upload to Google Sheets."""

from __future__ import annotations

import argparse
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


STATIC_SHEETS: dict[str, list[str]] = {
    "ShiftDefinitions": ["id", "name", "start_time", "end_time", "crosses_midnight", "sort_order", "is_active", "created_at", "updated_at", "version"],
    "SopCategories": ["id", "shift_definition_id", "name", "sort_order", "is_active", "created_at", "updated_at", "version", "color"],
    "ChecklistPoints": ["id", "sop_category_id", "title", "instruction", "input_type", "input_types", "options", "is_required", "target_time", "tolerance_minutes", "active_days", "number_min", "number_max", "sort_order", "is_active", "points", "created_at", "updated_at", "version"],
    "HandoverFields": ["id", "shift_definition_id", "label", "field_type", "input_types", "options", "is_required", "sort_order", "is_active", "created_at", "updated_at", "version"],
    "ShiftInstances": ["id", "shift_definition_id", "shift_date", "tab_month", "status", "pj_user_id", "opened_by", "opened_at", "opened_outside_hours", "closed_at", "closed_by", "close_type", "force_close_reason", "is_incomplete", "no_incident_confirmed", "void_reason", "void_by", "void_at", "is_test", "snapshot_encoding", "template_snapshot", "snapshot_hash", "created_at", "updated_at", "version", "tab_year"],
    "Participants": ["id", "shift_instance_id", "user_id", "first_action_at", "first_action_type", "created_at"],
    "Reports": ["id", "shift_instance_id", "report_number", "generated_by", "generated_at", "is_locked", "summary_stats", "content_hash", "unlock_count", "last_unlocked_at", "last_unlocked_by", "created_at", "updated_at", "version"],
    "Addenda": ["id", "report_id", "author_id", "note", "created_at"],
    "Summary": ["id", "summary_date", "shift_definition_id", "shifts_total", "shifts_closed_normal", "shifts_closed_forced", "shifts_void", "required_total", "required_done", "required_skipped", "timed_on_time", "timed_early", "timed_late", "incidents_total", "incidents_open", "incidents_by_category", "handovers_read", "participants_count", "computed_at"],
    "Snapshots": ["id", "shift_instance_id", "part_no", "chunk", "created_at"],
    "IncidentIndex": ["incident_id", "tab_month", "status", "category_id", "shift_instance_id", "outside_shift", "reported_at", "is_test", "updated_at"],
    "IncidentCategories": ["id", "name", "sort_order", "is_active", "points", "penalty_min_severity", "created_at", "updated_at"],
    "Notifications": ["id", "user_id", "type", "title", "body", "link", "is_read", "created_at"],
    "Drafts": ["id", "shift_instance_id", "user_id", "shift_date", "payload", "client_updated_at", "created_at", "updated_at", "version"],
    "ShiftPjAssignments": ["id", "shift_instance_id", "user_id", "started_at", "ended_at", "is_current", "source", "created_at"],
    "ScoreLedger": ["id", "scope_type", "subject_id", "branch_id", "shift_instance_id", "report_id", "metric_type", "score_version", "points_available", "points_earned", "incident_penalty", "final_points", "score_percent", "attribution_ratio", "calculated_at", "created_at"],
}

PERIOD_SHEETS: dict[str, list[str]] = {
    "Entries": ["id", "shift_instance_id", "point_ref", "state", "value", "out_of_range", "photo_ids", "completed_by", "completed_at", "timing_label", "timing_delta_minutes", "skip_reason", "created_at", "updated_at", "version"],
    "EntryLogs": ["id", "shift_instance_id", "entry_id", "point_ref", "action", "outcome", "user_id", "winner_user_id", "prev_state", "new_state", "value", "note", "client_action_id", "client_at", "at", "created_at"],
    "Handovers": ["id", "shift_instance_id", "values", "free_text", "photo_ids", "submitted_by", "submitted_at", "created_at", "updated_at"],
    "HandoverAcks": ["id", "handover_id", "reading_shift_instance_id", "user_id", "read_at", "created_at"],
    "Incidents": ["id", "shift_instance_id", "tab_month", "category_id", "description", "occurred_at", "reported_by", "reported_at", "status", "outside_shift", "link_source", "linked_by", "linked_at", "source_entry_id", "severity", "points_snapshot", "penalty_min_severity_snapshot", "status_changed_by", "status_changed_at", "is_test", "created_at", "updated_at", "version"],
    "IncidentNotes": ["id", "incident_id", "author_id", "author_role", "note", "created_at"],
    "Photos": ["id", "shift_instance_id", "owner_type", "owner_id", "storage", "file_ref", "mime", "size_bytes", "width", "height", "sort_order", "status", "uploaded_by", "uploaded_at", "purged_at", "created_at", "sha256"],
    "EntryReviews": ["id", "shift_instance_id", "entry_id", "point_ref", "worker_user_id", "review_status", "base_points", "earned_points", "awarded_points", "points_adjustment", "review_note", "reviewed_by", "reviewed_at", "review_stage", "version", "created_at"],
    "AuditLog": ["id", "seq", "at", "actor_id", "action", "object_type", "object_id", "branch_id", "shift_instance_id", "before", "after", "reason", "prev_hash", "hash"],
}

REGISTRY_SHEETS: dict[str, list[str]] = {
    "Daftar_Cabang": ["Cabang_ID", "Nama_Cabang", "Kode", "Timezone", "Spreadsheet_ID", "Folder_Drive_ID", "Aktif", "Alamat", "Arsip_Spreadsheet_IDs", "Branch_Folder_ID", "Evidence_Folder_ID"],
    "Users": ["User_ID", "Username", "PIN", "Nama", "Role", "Cabang_ID", "Aktif", "Must_Change_Pin", "Created_At"],
    "Settings_Global": ["Key", "Value"],
    "Share_Tokens": ["id", "report_id", "branch_id", "token", "created_by", "created_at", "expires_at", "revoked_at", "access_count", "last_accessed_at", "share_note", "visibility"],
    "Template_Referensi": ["Template_Spreadsheet_ID"],
}


def add_sheet(workbook: Workbook, title: str, headers: list[str]) -> None:
    sheet = workbook.create_sheet(title)
    sheet.append(headers)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
    sheet.row_dimensions[1].height = 28
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for index, header in enumerate(headers, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = min(max(len(header) + 2, 14), 30)


def generate_branch(output: Path, year: int) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    add_sheet(workbook, "_meta", ["key", "value"])
    meta = workbook["_meta"]
    meta.append(["schema_version", "2"])
    meta.append(["period_mode", "flat"])
    meta.append(["branch_id", ""])
    meta.append(["created_at", datetime.now(timezone.utc).isoformat()])
    meta.append(["tolerance_default_minutes", "15"])
    meta.append(["version", "2"])
    meta.column_dimensions["A"].width = 32
    meta.column_dimensions["B"].width = 42
    for title, headers in STATIC_SHEETS.items():
        add_sheet(workbook, title, headers)
    for base, headers in PERIOD_SHEETS.items():
        add_sheet(workbook, base, headers)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    print(f"created {output} ({len(workbook.sheetnames)} sheets, type=branch, mode=flat)")


def generate_registry(output: Path) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, headers in REGISTRY_SHEETS.items():
        add_sheet(workbook, title, headers)
    workbook["Users"].append([
        "TSDG8967",
        "admin",
        "123456",
        "admin",
        "admin",
        "",
        "TRUE",
        "YES",
        datetime.now(timezone.utc).isoformat(),
    ])
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    print(f"created {output} ({len(workbook.sheetnames)} sheets, type=registry)")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--kind", choices=["branch", "registry", "all"], default="all")
    parser.add_argument("--year", type=int, default=datetime.now(timezone.utc).year)
    parser.add_argument("--output", type=Path, default=Path("templates/mycheck-template-flat.xlsx"))
    args = parser.parse_args()
    if args.year < 2000 or args.year > 9999:
        raise SystemExit("--year harus berupa tahun 4 digit")

    if args.kind in ("branch", "all"):
        generate_branch(args.output, args.year)
    if args.kind in ("registry", "all"):
        registry_output = args.output.parent / "mycheck-registry.xlsx" if args.kind == "all" else args.output
        generate_registry(registry_output)


if __name__ == "__main__":
    main()
